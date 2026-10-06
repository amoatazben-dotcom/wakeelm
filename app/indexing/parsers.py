import ast
import io
import mimetypes
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath

from charset_normalizer import from_bytes

from app.core.exceptions import SafeError

LANGUAGES = {
    ".py": "Python",
    ".js": "JavaScript",
    ".jsx": "JavaScript",
    ".ts": "TypeScript",
    ".tsx": "TypeScript",
    ".java": "Java",
    ".kt": "Kotlin",
    ".kts": "Kotlin",
    ".dart": "Dart",
    ".go": "Go",
    ".rs": "Rust",
    ".c": "C",
    ".h": "C",
    ".cpp": "C++",
    ".hpp": "C++",
    ".cs": "C#",
    ".php": "PHP",
    ".rb": "Ruby",
    ".swift": "Swift",
    ".sh": "Shell",
    ".sql": "SQL",
    ".html": "HTML",
    ".css": "CSS",
    ".scss": "SCSS",
    ".vue": "Vue",
    ".svelte": "Svelte",
}
TEXT_EXTENSIONS = set(LANGUAGES) | {
    ".txt",
    ".md",
    ".rst",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".csv",
    ".toml",
    ".lock",
    ".mod",
    ".gradle",
    ".example",
    ".ini",
    ".cfg",
    ".gitignore",
    ".dockerignore",
    ".properties",
}
TEXT_NAMES = {
    "Dockerfile",
    "Makefile",
    "yarn.lock",
    "go.mod",
    ".gitignore",
    ".dockerignore",
    "LICENSE",
}


@dataclass
class ParsedFile:
    text: str | None
    encoding: str | None
    mime: str
    language: str | None
    symbols: list


def decode_text(data):
    for encoding, bom in [
        ("utf-8-sig", b"\xef\xbb\xbf"),
        ("utf-16", b"\xff\xfe"),
        ("utf-16", b"\xfe\xff"),
    ]:
        if data.startswith(bom):
            try:
                return data.decode(encoding), encoding
            except UnicodeError:
                return None, None
    if b"\x00" in data:
        return None, None
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeError:
        best = from_bytes(data).best()
        if best and best.chaos < 0.1 and best.coherence >= 0.5:
            try:
                return data.decode(best.encoding), best.encoding
            except UnicodeError:
                pass
    return None, None


class TextParser:
    def parse(self, data):
        return decode_text(data)


class MarkdownParser(TextParser):
    pass


class SourceCodeParser(TextParser):
    def symbols(self, text, language):
        if language != "Python":
            return []
        try:
            tree = ast.parse(text)
        except (SyntaxError, ValueError, RecursionError):
            return []
        output = []

        def walk(node, prefix=""):
            for child in ast.iter_child_nodes(node):
                if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    qualified = prefix + child.name
                    kind = (
                        "class"
                        if isinstance(child, ast.ClassDef)
                        else (
                            "test_function"
                            if child.name.startswith("test_")
                            else ("method" if prefix else "function")
                        )
                    )
                    output.append(
                        {
                            "symbol_type": kind,
                            "name": child.name[:255],
                            "qualified_name": qualified[:512],
                            "start_line": child.lineno,
                            "end_line": getattr(child, "end_lineno", child.lineno),
                            "signature": text.splitlines()[child.lineno - 1][:500],
                            "metadata_json": {"parser": "python_ast"},
                        }
                    )
                    walk(child, qualified + ".")
                elif isinstance(child, (ast.Import, ast.ImportFrom)):
                    for alias in child.names:
                        output.append(
                            {
                                "symbol_type": "import",
                                "name": alias.name[:255],
                                "qualified_name": alias.name[:512],
                                "start_line": child.lineno,
                                "end_line": child.end_lineno,
                                "signature": ast.unparse(child)[:500],
                                "metadata_json": {"parser": "python_ast"},
                            }
                        )
                elif isinstance(child, ast.Assign):
                    for target in child.targets:
                        if isinstance(target, ast.Name) and target.id.isupper():
                            output.append(
                                {
                                    "symbol_type": "constant",
                                    "name": target.id,
                                    "qualified_name": prefix + target.id,
                                    "start_line": child.lineno,
                                    "end_line": child.end_lineno,
                                    "signature": text.splitlines()[child.lineno - 1][:500],
                                    "metadata_json": {"parser": "python_ast"},
                                }
                            )
                else:
                    walk(child, prefix)

        walk(tree)
        return output[:5000]


class PDFParser:
    def parse(self, data):
        from pypdf import PdfReader

        if not data.startswith(b"%PDF-"):
            raise SafeError("INVALID_DOCUMENT")
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise SafeError("INVALID_DOCUMENT")
        return "\n".join(
            (page.extract_text() or "")[:100000]
            for page in __import__("itertools").islice(reader.pages, 100)
        ), "extracted"


class DocxParser:
    def parse(self, data):
        from defusedxml import ElementTree

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
        tree = ElementTree.fromstring(xml)
        return "\n".join(
            node.text or "" for node in tree.iter() if node.tag.endswith("}t")
        ), "extracted"


class SpreadsheetParser:
    def parse(self, data):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        rows = []
        try:
            for sheet in book.worksheets[:20]:
                rows.append(sheet.title)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row or 0, 5000),
                    max_col=min(sheet.max_column or 0, 100),
                    values_only=True,
                ):
                    rows.append(
                        "\t".join(str(cell)[:1000] if cell is not None else "" for cell in row)
                    )
                    if sum(len(x) for x in rows) > 2_000_000:
                        raise SafeError("FILE_TOO_LARGE")
        finally:
            book.close()
        return "\n".join(rows), "extracted"


class ParserRegistry:
    def __init__(self, settings):
        self.settings = settings

    def parse(self, path, data):
        file = PurePosixPath(path)
        extension = file.suffix.lower()
        mime = mimetypes.guess_type(path)[0] or "application/octet-stream"
        language = LANGUAGES.get(extension)
        parser = None
        if extension in {".docx", ".xlsx"}:
            self._check_document_zip(data)
            parser = DocxParser() if extension == ".docx" else SpreadsheetParser()
        elif extension == ".pdf":
            parser = PDFParser()
        elif extension in TEXT_EXTENSIONS or file.name in TEXT_NAMES:
            parser = (
                SourceCodeParser()
                if language
                else (MarkdownParser() if extension == ".md" else TextParser())
            )
        if parser is None:
            return ParsedFile(None, None, mime, language, [])
        try:
            text, encoding = parser.parse(data)
            if (
                text is not None
                and len(text.encode("utf-8")) > self.settings.max_single_file_size_mb * 1024**2
            ):
                raise SafeError("FILE_TOO_LARGE")
            symbols = SourceCodeParser().symbols(text, language) if text and language else []
            return ParsedFile(text, encoding, mime, language, symbols)
        except SafeError:
            raise
        except Exception:
            # Malformed document/parser failures cannot abort unrelated project indexing.
            return ParsedFile(None, None, mime, language, [])

    def _check_document_zip(self, data):
        if not data.startswith(b"PK"):
            raise SafeError("INVALID_DOCUMENT")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > self.settings.max_archive_files
                    or sum(m.file_size for m in members)
                    > self.settings.max_extracted_size_mb * 1024**2
                ):
                    raise SafeError("EXTRACTED_TOO_LARGE")
                if any(
                    m.file_size > self.settings.max_single_file_size_mb * 1024**2
                    or m.file_size > max(1024**2, m.compress_size * 1000)
                    for m in members
                ):
                    raise SafeError("FILE_TOO_LARGE")
        except zipfile.BadZipFile:
            raise SafeError("INVALID_DOCUMENT") from None
